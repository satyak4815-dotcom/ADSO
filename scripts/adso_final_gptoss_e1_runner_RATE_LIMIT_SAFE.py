import os
import sys
import json
import time
import re
from datetime import datetime, timezone
from pathlib import Path

from groq import Groq
from openpyxl import load_workbook

MODEL = "openai/gpt-oss-120b"
SHEET = "Benchmark Instances"
MAX_RETRIES = 8
RETRY_DELAY_SECONDS = 2

SYSTEM_MESSAGE = """You are generating diagnostic evidence for a code-efficiency experiment.

Return only one valid JSON object with exactly these fields:
"time_complexity": string,
"auxiliary_space_complexity": string,
"main_bottleneck": string,
"avoidable_work": string,
"constraint_fit": string,
"optimality_rationale": string,
"diagnostic_confidence": integer from 0 to 100.

Rules:
- Do NOT output OPTIMIZE, PRESERVE, ESCALATE, or any equivalent intervention decision.
- Do NOT assume access to another implementation, hidden labels, profiler results, hidden tests, or external tools.
- Judge only from the problem statement, constraints, required output, and source code.
- Keep each field concise and technical.
- Do not include markdown fences or text outside the JSON object.
"""

REQUIRED_FIELDS = {
    "time_complexity",
    "auxiliary_space_complexity",
    "main_bottleneck",
    "avoidable_work",
    "constraint_fit",
    "optimality_rationale",
    "diagnostic_confidence",
}

def validate_payload(payload):
    if not isinstance(payload, dict):
        return False, "Response is not a JSON object"
    if set(payload.keys()) != REQUIRED_FIELDS:
        return False, f"Unexpected fields: {sorted(payload.keys())}"
    c = payload["diagnostic_confidence"]
    if isinstance(c, bool) or not isinstance(c, int) or not 0 <= c <= 100:
        return False, "diagnostic_confidence must be integer 0..100"
    for k in REQUIRED_FIELDS - {"diagnostic_confidence"}:
        if not isinstance(payload[k], str):
            return False, f"{k} must be a string"
    return True, None

def load_existing(output_path):
    completed = set()
    p = Path(output_path)
    if not p.exists():
        return completed
    with p.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
                if row.get("status") == "SUCCESS":
                    completed.add(row.get("instance_id"))
            except Exception:
                pass
    return completed

def build_prompt(row, h):
    return f"""You are analyzing a functionally correct Python program for efficiency characteristics.

Problem:
{row[h["Problem"]]}

Input and constraints:
{row[h["Input/Constraints"]]}

Required output:
{row[h["Output"]]}

Source code:
```python
{row[h["Source Code"]]}
```

Produce diagnostic evidence only. Do not make an OPTIMIZE/PRESERVE/ESCALATE decision."""

def main():
    if len(sys.argv) != 3:
        print("Usage: python adso_final_gptoss_e1_runner.py INPUT_XLSX OUTPUT_JSONL")
        sys.exit(1)

    input_xlsx, output_jsonl = sys.argv[1], sys.argv[2]

    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        print("ERROR: GROQ_API_KEY is not set.")
        print('PowerShell: $env:GROQ_API_KEY="YOUR_KEY"')
        sys.exit(1)

    wb = load_workbook(input_xlsx, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f'ERROR: "{SHEET}" sheet not found.')
        sys.exit(1)

    ws = wb[SHEET]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    h = {str(v): i for i, v in enumerate(headers) if v is not None}

    required_cols = ["Instance ID", "Problem", "Input/Constraints", "Output", "Source Code"]
    for col in required_cols:
        if col not in h:
            print(f'ERROR: Required column "{col}" not found.')
            sys.exit(1)

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[h["Instance ID"]]:
            continue
        rows.append((str(row[h["Instance ID"]]), build_prompt(row, h)))

    completed = load_existing(output_jsonl)
    client = Groq(api_key=api_key)

    print(f"Model: {MODEL}")
    print(f"Loaded {len(rows)} held-out instances.")
    print(f"Already completed successfully: {len(completed)}")
    print("Starting/resuming E1 generation...\n")

    with open(output_jsonl, "a", encoding="utf-8") as out:
        for idx, (instance_id, prompt) in enumerate(rows, start=1):
            if instance_id in completed:
                print(f"[{idx:03d}/{len(rows)}] SKIP    {instance_id}")
                continue

            record = None

            for attempt in range(1, MAX_RETRIES + 1):
                started = time.perf_counter()
                try:
                    response = client.chat.completions.create(
                        model=MODEL,
                        messages=[
                            {"role": "system", "content": SYSTEM_MESSAGE},
                            {"role": "user", "content": prompt},
                        ],
                        response_format={"type": "json_object"},
                    )

                    latency = time.perf_counter() - started
                    raw = response.choices[0].message.content or ""

                    try:
                        payload = json.loads(raw)
                    except Exception as e:
                        record = {
                            "instance_id": instance_id,
                            "model": MODEL,
                            "provider": "Groq",
                            "stage": "E1",
                            "status": "MALFORMED",
                            "raw_response": raw,
                            "error": f"JSON parse failure: {e}",
                            "latency_seconds": round(latency, 4),
                            "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                        }
                        break

                    valid, error = validate_payload(payload)
                    if not valid:
                        record = {
                            "instance_id": instance_id,
                            "model": MODEL,
                            "provider": "Groq",
                            "stage": "E1",
                            "status": "MALFORMED",
                            "raw_response": raw,
                            "error": error,
                            "latency_seconds": round(latency, 4),
                            "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                        }
                        break

                    usage = getattr(response, "usage", None)
                    record = {
                        "instance_id": instance_id,
                        "model": MODEL,
                        "provider": "Groq",
                        "stage": "E1",
                        "status": "SUCCESS",
                        **payload,
                        "raw_response": raw,
                        "latency_seconds": round(latency, 4),
                        "prompt_tokens": getattr(usage, "prompt_tokens", None) if usage else None,
                        "completion_tokens": getattr(usage, "completion_tokens", None) if usage else None,
                        "total_tokens": getattr(usage, "total_tokens", None) if usage else None,
                        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                    }
                    break

                except Exception as e:
                    latency = time.perf_counter() - started
                    msg = str(e)

                    # Groq 429 responses often include "Please try again in 10m13.44s".
                    # Wait for the requested reset instead of repeatedly burning retries.
                    wait_seconds = None
                    if "429" in msg or "rate_limit" in msg.lower() or "rate limit" in msg.lower():
                        m = re.search(r"try again in\s+(?:(\d+)m)?([0-9.]+)s", msg, re.I)
                        if m:
                            minutes = int(m.group(1) or 0)
                            seconds = float(m.group(2) or 0)
                            wait_seconds = minutes * 60 + seconds + 5

                    if attempt < MAX_RETRIES:
                        if wait_seconds is not None:
                            print(
                                f"[{idx:03d}/{len(rows)}] RATE_LIMIT {instance_id} "
                                f"attempt {attempt}; sleeping {wait_seconds:.1f}s until quota window resets."
                            )
                            time.sleep(wait_seconds)
                        else:
                            print(f"[{idx:03d}/{len(rows)}] RETRY   {instance_id} attempt {attempt}: {e}")
                            time.sleep(RETRY_DELAY_SECONDS)
                        continue

                    record = {
                        "instance_id": instance_id,
                        "model": MODEL,
                        "provider": "Groq",
                        "stage": "E1",
                        "status": "API_ERROR",
                        "attempt": attempt,
                        "error": str(e),
                        "latency_seconds": round(latency, 4),
                        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                    }

            out.write(json.dumps(record, ensure_ascii=False) + "\n")
            out.flush()
            print(f"[{idx:03d}/{len(rows)}] {record['status']:<9} {instance_id}")

    print(f"\nFinished. Results saved to: {output_jsonl}")
    print("Do not manually edit model responses. Upload the JSONL file to ChatGPT.")

if __name__ == "__main__":
    main()
