# ADSO Final Held-out Llama-3.1-8B C0-C3 Runner v2
# Deterministic first-JSON extraction + migration of prior preflight generations.
#
# Scientific handling:
# - Prompts/model settings are unchanged from v1.
# - Full raw_response is always preserved.
# - Parser accepts, in order:
#     1) direct JSON
#     2) first ```json ... ``` fenced object, even if commentary follows
# - Extracted JSON must have EXACTLY: decision, confidence, reason.
# - Existing v1 generations are REPARSED, not regenerated.
# - Original v1 file is never modified.
# - PRIVATE Ground Truth is never read.
# - ADSO_TEST_050 E1 raw response is preserved verbatim; no repair/selective rerun.
#
# Run in Colab with %run -i (model/tokenizer must already be loaded).

import argparse, json, re, sys, time
from datetime import datetime, timezone
from pathlib import Path
import openpyxl, torch

MODEL_NAME = "meta-llama/Meta-Llama-3.1-8B-Instruct"
CONDITIONS = ("C0", "C1", "C2", "C3")
MAX_NEW_TOKENS = 220

OUTPUT_SCHEMA = """Return one JSON object with exactly these fields:
"decision": "OPTIMIZE" or "PRESERVE",
"confidence": integer from 0 to 100,
"reason": string.

Judge only from the information above. Do not assume access to another implementation, benchmark label, profiler result, hidden ground truth, or future evidence."""

def now_utc():
    return datetime.now(timezone.utc).isoformat()

def sheet_map(wb, sheet, col):
    rows = wb[sheet].iter_rows(values_only=True)
    h = list(next(rows))
    ii, vi = h.index("Instance ID"), h.index(col)
    return {str(r[ii]): "" if r[vi] is None else str(r[vi])
            for r in rows if r and r[ii]}

def load_e1(path):
    d = {}
    with open(path, encoding="utf-8") as f:
        for n, line in enumerate(f, 1):
            if not line.strip():
                continue
            x = json.loads(line)
            iid = x.get("instance_id")
            if not iid or iid in d:
                raise RuntimeError(f"E1 ID problem at line {n}: {iid}")
            raw = x.get("raw_response")
            if raw is None or not str(raw).strip():
                raise RuntimeError(f"Missing raw E1: {iid}")
            d[iid] = {
                "raw": str(raw),
                "status": x.get("status"),
                "parse_method": x.get("parse_method")
            }
    return d

def build_prompt(e0, e1, e2, e3, c):
    if c == "C0":
        return e0

    marker = "Return one JSON object with exactly these fields:"
    if marker not in e0:
        raise RuntimeError("E0 schema marker missing")

    body = e0.split(marker, 1)[0].rstrip()
    p = (
        body
        + "\n\nE1 — Model-generated structured self-diagnosis from an earlier independent pass "
          "over the same E0 information:\n"
        + e1
        + "\n\nTreat E1 as a model-generated deliberation signal, not as independent external evidence."
    )
    if c in ("C2", "C3"):
        p += "\n\nE2 — Deterministic static evidence:\n" + e2
    if c == "C3":
        p += "\n\nE3 — Controlled dynamic evidence:\n" + e3
    return p + "\n\n" + OUTPUT_SCHEMA

def validate_obj(x):
    if not isinstance(x, dict):
        return None, "Output is not a JSON object"
    required = {"decision", "confidence", "reason"}
    if set(x.keys()) != required:
        return None, f"Fields must be exactly {sorted(required)}; got {sorted(x.keys())}"
    if x["decision"] not in ("OPTIMIZE", "PRESERVE"):
        return None, f"Invalid decision: {x['decision']!r}"
    c = x["confidence"]
    if isinstance(c, bool) or not isinstance(c, int) or not 0 <= c <= 100:
        return None, f"Invalid confidence: {c!r}"
    if not isinstance(x["reason"], str):
        return None, "reason must be a string"
    return x, None

def parse_output(text):
    """Deterministic parser. Does not infer/repair JSON."""
    s = (text or "").strip()

    # Rule 1: whole response is JSON.
    try:
        x = json.loads(s)
        x, err = validate_obj(x)
        if x:
            return x, "direct_json", None
        return None, None, err
    except Exception:
        pass

    # Rule 2: first explicitly fenced JSON object.
    # Commentary before/after the fence is ignored, but JSON inside is not repaired.
    m = re.search(r"```json\s*(\{.*?\})\s*```", s, flags=re.I | re.S)
    if m:
        try:
            x = json.loads(m.group(1))
        except Exception as e:
            return None, None, f"First fenced JSON parse failure: {e}"
        x, err = validate_obj(x)
        if x:
            return x, "first_json_fence_extracted", None
        return None, None, err

    return None, None, "No valid direct JSON or fenced JSON object"

def read_records(path):
    records = []
    p = Path(path)
    if not p.exists():
        return records
    with p.open(encoding="utf-8") as f:
        for n, line in enumerate(f, 1):
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except Exception as e:
                raise RuntimeError(f"Invalid JSONL at {path}, line {n}: {e}")
    return records

def migrate_existing(old_path, new_path):
    """Copy/reparse old generations into new v2 output. Never changes old file."""
    if not old_path or not Path(old_path).exists():
        return 0

    if Path(new_path).exists() and Path(new_path).stat().st_size > 0:
        # Migration is only for a fresh v2 output. Existing v2 output resumes normally.
        return 0

    old = read_records(old_path)
    migrated = 0
    seen = set()

    with open(new_path, "a", encoding="utf-8") as out:
        for r in old:
            iid, cond = r.get("instance_id"), r.get("condition")
            key = (iid, cond)
            if not iid or cond not in CONDITIONS or key in seen:
                continue
            raw = r.get("raw_response")
            if raw is None:
                continue

            parsed, method, err = parse_output(raw)
            nr = dict(r)
            nr["original_status_v1"] = r.get("status")
            nr["original_parse_method_v1"] = r.get("parse_method")
            nr["parser_version"] = "v2_first_json_fence"
            nr["migration_source"] = str(old_path)
            nr["migrated_not_regenerated"] = True
            nr["status"] = "SUCCESS" if parsed else "MALFORMED"
            nr["parse_method"] = method
            nr["decision"] = parsed["decision"] if parsed else None
            nr["confidence"] = parsed["confidence"] if parsed else None
            nr["reason"] = parsed["reason"] if parsed else None
            nr["error"] = err
            nr["v2_migration_timestamp_utc"] = now_utc()

            out.write(json.dumps(nr, ensure_ascii=False) + "\n")
            out.flush()
            seen.add(key)
            migrated += 1

    return migrated

def completed(path):
    s = set()
    for x in read_records(path):
        if x.get("status") in ("SUCCESS", "MALFORMED"):
            iid, c = x.get("instance_id"), x.get("condition")
            if iid and c:
                s.add((iid, c))
    return s

def generate(prompt):
    msgs = [{"role": "user", "content": prompt}]
    rendered = tokenizer.apply_chat_template(
        msgs, tokenize=False, add_generation_prompt=True
    )
    inp = tokenizer(rendered, return_tensors="pt")
    dev = next(model.parameters()).device
    inp = {k: v.to(dev) for k, v in inp.items()}
    nin = int(inp["input_ids"].shape[-1])

    t0 = time.time()
    with torch.inference_mode():
        out = model.generate(
            **inp,
            max_new_tokens=MAX_NEW_TOKENS,
            do_sample=False,
            pad_token_id=tokenizer.eos_token_id
        )
    latency = time.time() - t0
    new = out[0][inp["input_ids"].shape[-1]:]
    raw = tokenizer.decode(new, skip_special_tokens=True)
    return raw, latency, nin, int(new.shape[-1])

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--e1", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--migrate-from", default=None,
                    help="Optional old v1 JSONL. Its raw generations are reparsed, never regenerated.")
    a = ap.parse_args()

    if "model" not in globals() or "tokenizer" not in globals():
        print("ERROR: model/tokenizer not loaded in this Python process. In Colab use %run -i.")
        sys.exit(2)

    wb = openpyxl.load_workbook(a.workbook, read_only=True, data_only=True)
    for s in ("E0 Prompts", "E2 Static Evidence", "E3 Dynamic Evidence"):
        if s not in wb.sheetnames:
            raise RuntimeError(f"Missing sheet: {s}")

    # Intentionally NEVER access PRIVATE Ground Truth.
    e0 = sheet_map(wb, "E0 Prompts", "Frozen E0 Decision Prompt")
    e2 = sheet_map(wb, "E2 Static Evidence", "Neutral E2 Evidence Text")
    e3 = sheet_map(wb, "E3 Dynamic Evidence", "Neutral E3 Evidence Text")
    e1 = load_e1(a.e1)

    ids = list(e0)
    if len(ids) != 120:
        raise RuntimeError(f"Expected 120 instances, found {len(ids)}")
    for iid in ids:
        if iid not in e1 or iid not in e2 or iid not in e3:
            raise RuntimeError(f"Evidence missing for {iid}")
    if set(e1) != set(ids):
        raise RuntimeError("E1 IDs do not exactly match frozen benchmark IDs")

    Path(a.output).parent.mkdir(parents=True, exist_ok=True)

    migrated = migrate_existing(a.migrate_from, a.output)
    runs = [(iid, c) for iid in ids for c in CONDITIONS]
    done = completed(a.output)

    print(f"Model: {MODEL_NAME}")
    print("Parser: v2 deterministic first fenced JSON extraction")
    print(f"Loaded {len(runs)} runs (120 instances x 4 conditions).")
    if a.migrate_from:
        print(f"Migrated/reparsed existing v1 generations: {migrated}")
    print(f"Already completed in v2 output: {sum(k in done for k in runs)}")
    print(f"E1: {sum(v['status']=='SUCCESS' for v in e1.values())} SUCCESS, "
          f"{sum(v['status']=='MALFORMED' for v in e1.values())} MALFORMED")
    print("TEST_050 raw E1 preserved verbatim; no repair/selective rerun.")
    print("PRIVATE Ground Truth is NOT loaded.\n")

    with open(a.output, "a", encoding="utf-8") as f:
        for n, (iid, c) in enumerate(runs, 1):
            if (iid, c) in done:
                continue

            prompt = build_prompt(e0[iid], e1[iid]["raw"], e2[iid], e3[iid], c)
            print(f"[{n:03d}/480] {iid} {c} running...", flush=True)

            try:
                raw, lat, nin, nout = generate(prompt)
                x, method, err = parse_output(raw)

                rec = {
                    "instance_id": iid,
                    "condition": c,
                    "model": MODEL_NAME,
                    "provider": "Hugging Face / local Colab",
                    "stage": "C0-C3 final heldout decision",
                    "status": "SUCCESS" if x else "MALFORMED",
                    "parser_version": "v2_first_json_fence",
                    "parse_method": method,
                    "decision": x["decision"] if x else None,
                    "confidence": x["confidence"] if x else None,
                    "reason": x["reason"] if x else None,
                    "raw_response": raw,
                    "error": err,
                    "e1_source_status": e1[iid]["status"] if c != "C0" else None,
                    "e1_source_parse_method": e1[iid]["parse_method"] if c != "C0" else None,
                    "latency_seconds": round(lat, 4),
                    "input_tokens": nin,
                    "output_tokens": nout,
                    "do_sample": False,
                    "max_new_tokens": MAX_NEW_TOKENS,
                    "timestamp_utc": now_utc()
                }
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                f.flush()

                if x:
                    print(f"          SUCCESS {x['decision']} confidence={x['confidence']} ({method})", flush=True)
                else:
                    print(f"          MALFORMED preserved: {err}", flush=True)

            except KeyboardInterrupt:
                print("\nStopped. v2 checkpoint safe; rerun the same %run -i command to resume.")
                raise
            except Exception as ex:
                rec = {
                    "instance_id": iid,
                    "condition": c,
                    "model": MODEL_NAME,
                    "status": "ERROR",
                    "parser_version": "v2_first_json_fence",
                    "error": repr(ex),
                    "timestamp_utc": now_utc()
                }
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                f.flush()
                print(f"          ERROR recorded: {ex!r}", flush=True)

    print(f"\nCheckpoint: {sum(k in completed(a.output) for k in runs)}/480 completed")
    print("Output:", a.output)

if __name__ == "__main__":
    main()
