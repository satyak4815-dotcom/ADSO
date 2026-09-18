# ADSO Final Held-out E1 Runner — Meta Llama 3.1 8B Instruct
# Intended for Google Colab after `model` and `tokenizer` are already loaded.
# Saves incrementally and resumes by skipping prior SUCCESS rows.

import json
import re
import time
from pathlib import Path
from datetime import datetime, timezone

import torch
from openpyxl import load_workbook

MODEL_ID = "meta-llama/Meta-Llama-3.1-8B-Instruct"
STAGE = "E1"
MAX_NEW_TOKENS = 320
DO_SAMPLE = False

SYSTEM_TEXT = """You are generating diagnostic evidence for a code-efficiency experiment.

Return only one valid JSON object with exactly these fields:
"time_complexity": string,
"auxiliary_space_complexity": string,
"main_bottleneck": string,
"avoidable_work": string,
"constraint_fit": string,
"optimality_rationale": string,
"diagnostic_confidence": integer from 0 to 100.

Rules:
- Do NOT output OPTIMIZE, PRESERVE, ESCALATE, or any equivalent intervention decision.
- Do NOT assume access to another implementation, hidden labels, profiler results, hidden tests, or external tools.
- Judge only from the problem statement, constraints, required output, and source code.
- Keep each field concise and technical.
- Do not include explanatory text outside the JSON object.
"""

REQUIRED_FIELDS = {
    "time_complexity",
    "auxiliary_space_complexity",
    "main_bottleneck",
    "avoidable_work",
    "constraint_fit",
    "optimality_rationale",
    "diagnostic_confidence",
}

def parse_response(raw):
    """
    Frozen parser policy:
    1) try direct JSON;
    2) if that fails, accept exactly one markdown fenced JSON object;
    3) otherwise MALFORMED.
    No semantic repair is performed. Raw text is always preserved.
    """
    try:
        return json.loads(raw), "direct_json", None
    except Exception as e1:
        m = re.fullmatch(r"\s*```(?:json)?\s*(\{.*\})\s*```\s*", raw, flags=re.I | re.S)
        if m:
            try:
                return json.loads(m.group(1)), "single_json_fence_removed", None
            except Exception as e2:
                return None, None, f"Fenced JSON parse failure: {e2}"
        return None, None, f"Direct JSON parse failure: {e1}"

def validate_payload(payload):
    if not isinstance(payload, dict):
        return False, "Response is not a JSON object"
    if set(payload.keys()) != REQUIRED_FIELDS:
        return False, f"Unexpected fields: {sorted(payload.keys())}"
    c = payload["diagnostic_confidence"]
    if isinstance(c, bool) or not isinstance(c, int) or not 0 <= c <= 100:
        return False, "diagnostic_confidence must be an integer from 0 to 100"
    for k in REQUIRED_FIELDS - {"diagnostic_confidence"}:
        if not isinstance(payload[k], str):
            return False, f"{k} must be a string"
    return True, None

def completed_ids(output_jsonl):
    done = set()
    p = Path(output_jsonl)
    if not p.exists():
        return done
    with p.open("r", encoding="utf-8") as f:
        for line in f:
            try:
                row = json.loads(line)
                if row.get("status") == "SUCCESS":
                    done.add(row.get("instance_id"))
            except Exception:
                pass
    return done

def build_user_prompt(problem, inp, out, code):
    return f"""You are analyzing a functionally correct Python program for efficiency characteristics.

Problem:
{problem}

Input and constraints:
{inp}

Required output:
{out}

Source code:
```python
{code}
```

Produce diagnostic evidence only. Do not make an OPTIMIZE/PRESERVE/ESCALATE decision."""

def run_adso_llama_e1(
    benchmark_xlsx="/content/ADSO_Final_Heldout_Benchmark_v1.0_FROZEN.xlsx",
    output_jsonl="/content/ADSO_Final_Llama31_8B_E1_Results.jsonl",
):
    # Require the already-loaded objects from previous Colab cells.
    if "model" not in globals() or "tokenizer" not in globals():
        raise RuntimeError("`model` and `tokenizer` are not loaded in this Colab runtime.")

    model.eval()

    wb = load_workbook(benchmark_xlsx, read_only=True, data_only=True)
    ws = wb["Benchmark Instances"]

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    h = {v: i for i, v in enumerate(header)}
    needed = ["Instance ID", "Problem", "Input/Constraints", "Output", "Source Code"]
    for col in needed:
        if col not in h:
            raise RuntimeError(f"Missing required column: {col}")

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        iid = row[h["Instance ID"]]
        if iid:
            rows.append({
                "instance_id": str(iid),
                "prompt": build_user_prompt(
                    row[h["Problem"]],
                    row[h["Input/Constraints"]],
                    row[h["Output"]],
                    row[h["Source Code"]],
                ),
            })

    done = completed_ids(output_jsonl)

    # Capture reproducibility details available from the loaded model/tokenizer.
    model_commit = getattr(getattr(model, "config", None), "_commit_hash", None)
    tokenizer_commit = getattr(tokenizer, "_commit_hash", None)

    print(f"Model: {MODEL_ID}")
    print(f"Instances: {len(rows)}")
    print(f"Already successful: {len(done)}")
    print(f"Model commit: {model_commit}")
    print(f"Tokenizer commit: {tokenizer_commit}")
    print(f"Generation: do_sample={DO_SAMPLE}, max_new_tokens={MAX_NEW_TOKENS}")
    print()

    with open(output_jsonl, "a", encoding="utf-8") as f:
        for n, item in enumerate(rows, 1):
            iid = item["instance_id"]
            if iid in done:
                print(f"[{n:03d}/{len(rows)}] SKIP      {iid}")
                continue

            messages = [
                {"role": "system", "content": SYSTEM_TEXT},
                {"role": "user", "content": item["prompt"]},
            ]

            try:
                formatted = tokenizer.apply_chat_template(
                    messages,
                    tokenize=False,
                    add_generation_prompt=True,
                )
                inputs = tokenizer(
                    formatted,
                    return_tensors="pt",
                    add_special_tokens=False,
                ).to(model.device)

                start = time.perf_counter()
                with torch.inference_mode():
                    outputs = model.generate(
                        **inputs,
                        max_new_tokens=MAX_NEW_TOKENS,
                        do_sample=DO_SAMPLE,
                        pad_token_id=tokenizer.eos_token_id,
                    )
                latency = time.perf_counter() - start

                new_tokens = outputs[0, inputs["input_ids"].shape[1]:]
                raw = tokenizer.decode(
                    new_tokens,
                    skip_special_tokens=True,
                    clean_up_tokenization_spaces=False,
                )

                payload, parse_method, parse_error = parse_response(raw)

                if payload is None:
                    record = {
                        "instance_id": iid,
                        "model": MODEL_ID,
                        "provider": "Hugging Face / local Colab",
                        "hardware": torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU",
                        "stage": STAGE,
                        "status": "MALFORMED",
                        "parse_method": None,
                        "raw_response": raw,
                        "error": parse_error,
                        "latency_seconds": round(latency, 4),
                        "input_tokens": int(inputs["input_ids"].shape[1]),
                        "output_tokens": int(new_tokens.shape[0]),
                        "do_sample": DO_SAMPLE,
                        "max_new_tokens": MAX_NEW_TOKENS,
                        "model_commit": model_commit,
                        "tokenizer_commit": tokenizer_commit,
                        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                    }
                else:
                    valid, err = validate_payload(payload)
                    if not valid:
                        record = {
                            "instance_id": iid,
                            "model": MODEL_ID,
                            "provider": "Hugging Face / local Colab",
                            "hardware": torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU",
                            "stage": STAGE,
                            "status": "MALFORMED",
                            "parse_method": parse_method,
                            "raw_response": raw,
                            "error": err,
                            "latency_seconds": round(latency, 4),
                            "input_tokens": int(inputs["input_ids"].shape[1]),
                            "output_tokens": int(new_tokens.shape[0]),
                            "do_sample": DO_SAMPLE,
                            "max_new_tokens": MAX_NEW_TOKENS,
                            "model_commit": model_commit,
                            "tokenizer_commit": tokenizer_commit,
                            "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                        }
                    else:
                        record = {
                            "instance_id": iid,
                            "model": MODEL_ID,
                            "provider": "Hugging Face / local Colab",
                            "hardware": torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU",
                            "stage": STAGE,
                            "status": "SUCCESS",
                            "parse_method": parse_method,
                            **payload,
                            "raw_response": raw,
                            "latency_seconds": round(latency, 4),
                            "input_tokens": int(inputs["input_ids"].shape[1]),
                            "output_tokens": int(new_tokens.shape[0]),
                            "do_sample": DO_SAMPLE,
                            "max_new_tokens": MAX_NEW_TOKENS,
                            "model_commit": model_commit,
                            "tokenizer_commit": tokenizer_commit,
                            "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                        }

            except Exception as e:
                record = {
                    "instance_id": iid,
                    "model": MODEL_ID,
                    "provider": "Hugging Face / local Colab",
                    "hardware": torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU",
                    "stage": STAGE,
                    "status": "TECHNICAL_FAILURE",
                    "error": repr(e),
                    "do_sample": DO_SAMPLE,
                    "max_new_tokens": MAX_NEW_TOKENS,
                    "model_commit": model_commit,
                    "tokenizer_commit": tokenizer_commit,
                    "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                }

            f.write(json.dumps(record, ensure_ascii=False) + "\n")
            f.flush()
            print(f"[{n:03d}/{len(rows)}] {record['status']:<17} {iid}")

    print()
    print("Finished/resumed run.")
    print(f"Saved to: {output_jsonl}")
    print("Upload the JSONL file to ChatGPT when complete.")
