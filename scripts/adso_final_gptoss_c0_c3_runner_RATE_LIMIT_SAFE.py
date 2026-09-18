import os, sys, json, time, re
from datetime import datetime, timezone
from pathlib import Path
from groq import Groq
from openpyxl import load_workbook

MODEL = "openai/gpt-oss-120b"
CONDITIONS = ("C0", "C1", "C2", "C3")
MAX_RETRIES = 20
TECHNICAL_RETRY_SECONDS = 5
SYSTEM_MESSAGE = '''You are making a code-efficiency intervention decision for a controlled experiment.
Return only one valid JSON object with exactly these fields:
"decision": "OPTIMIZE" or "PRESERVE",
"confidence": integer from 0 to 100,
"reason": string.
Rules:
- Judge only from the evidence supplied in this prompt.
- Do not assume access to another implementation, benchmark label, hidden ground truth, or evidence from a later condition.
- OPTIMIZE means efficiency-oriented modification is warranted.
- PRESERVE means efficiency-oriented modification is not warranted.
- Do not include markdown fences or text outside the JSON object.'''
REQUIRED_FIELDS = {"decision", "confidence", "reason"}
DECISION_TAIL = '''Return one JSON object with exactly these fields:
"decision": "OPTIMIZE" or "PRESERVE",
"confidence": integer from 0 to 100,
"reason": string.

Judge only from the information above. Do not assume access to another implementation, benchmark label, profiler result, hidden ground truth, or future evidence.'''

def validate_payload(p):
    if not isinstance(p, dict): return False, "Response is not a JSON object"
    if set(p) != REQUIRED_FIELDS: return False, f"Unexpected fields: {sorted(p)}"
    if p["decision"] not in {"OPTIMIZE","PRESERVE"}: return False, "invalid decision"
    c=p["confidence"]
    if isinstance(c,bool) or not isinstance(c,int) or not 0<=c<=100: return False,"invalid confidence"
    if not isinstance(p["reason"],str): return False,"reason must be string"
    return True,None

def load_successful_e1(path):
    d={}
    with open(path,encoding="utf-8") as f:
        for line in f:
            if not line.strip(): continue
            r=json.loads(line)
            if r.get("status")=="SUCCESS":
                d[r["instance_id"]]={k:r[k] for k in ["time_complexity","auxiliary_space_complexity","main_bottleneck","avoidable_work","constraint_fit","optimality_rationale","diagnostic_confidence"]}
    return d

def sheetmap(wb,sheet,valcol):
    ws=wb[sheet]; headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]; h={str(v):i for i,v in enumerate(headers) if v is not None}
    return {str(r[h["Instance ID"]]):str(r[h[valcol]] or "") for r in ws.iter_rows(min_row=2,values_only=True) if r[h["Instance ID"]]}

def base_e0(prompt):
    marker="Return one JSON object with exactly these fields:"
    if marker not in prompt: raise RuntimeError("Unexpected E0 prompt format")
    return prompt.split(marker,1)[0].rstrip()

def make_prompt(base,cond,e1,e2,e3):
    parts=[base]
    if cond in {"C1","C2","C3"}: parts.append("Structured self-diagnosis evidence (E1; generated previously by this same model):\n"+json.dumps(e1,ensure_ascii=False,indent=2))
    if cond in {"C2","C3"}: parts.append("Deterministic static evidence (E2):\n"+e2)
    if cond=="C3": parts.append("Controlled dynamic evidence (E3):\n"+e3)
    parts.append(DECISION_TAIL)
    return "\n\n".join(parts)

def completed(path):
    s=set(); p=Path(path)
    if not p.exists(): return s
    for line in p.open(encoding="utf-8"):
        try:
            r=json.loads(line)
            if r.get("status")=="SUCCESS": s.add((r.get("condition"),r.get("instance_id")))
        except: pass
    return s

def retry_after(msg):
    m=re.search(r"try again in\s+(?:(\d+)m)?([0-9.]+)s",msg,re.I)
    if m:return int(m.group(1) or 0)*60+float(m.group(2) or 0)
    return None

def main():
    if len(sys.argv)!=4:
        print("Usage: python adso_final_gptoss_c0_c3_runner_RATE_LIMIT_SAFE.py BENCHMARK.xlsx GPTOSS_E1.jsonl OUTPUT.jsonl");sys.exit(1)
    bench,e1file,outfile=sys.argv[1:4]
    key=os.getenv("GROQ_API_KEY")
    if not key: print("ERROR: GROQ_API_KEY is not set.");sys.exit(1)
    wb=load_workbook(bench,read_only=True,data_only=True)
    e0=sheetmap(wb,"E0 Prompts","Frozen E0 Decision Prompt")
    e2=sheetmap(wb,"E2 Static Evidence","Neutral E2 Evidence Text")
    e3=sheetmap(wb,"E3 Dynamic Evidence","Neutral E3 Evidence Text")
    e1=load_successful_e1(e1file)
    ids=list(e0)
    if len(ids)!=120: raise RuntimeError(f"Expected 120 instances, got {len(ids)}")
    miss=[x for x in ids if x not in e1 or x not in e2 or x not in e3]
    if miss: raise RuntimeError(f"Missing evidence for: {miss[:10]}")
    jobs=[(c,i,make_prompt(base_e0(e0[i]),c,e1[i],e2[i],e3[i])) for c in CONDITIONS for i in ids]
    done=completed(outfile); client=Groq(api_key=key)
    print(f"Model: {MODEL}\nJobs: 480 = 120 x 4\nAlready SUCCESS: {len(done)}\nOrder: C0 -> C1 -> C2 -> C3\n")
    with open(outfile,"a",encoding="utf-8") as out:
      for idx,(cond,iid,prompt) in enumerate(jobs,1):
        if (cond,iid) in done: print(f"[{idx:03d}/480] SKIP {cond} {iid}");continue
        rec=None
        for attempt in range(1,MAX_RETRIES+1):
          st=time.perf_counter()
          try:
            resp=client.chat.completions.create(model=MODEL,messages=[{"role":"system","content":SYSTEM_MESSAGE},{"role":"user","content":prompt}],response_format={"type":"json_object"})
            lat=time.perf_counter()-st; raw=resp.choices[0].message.content or ""
            try: payload=json.loads(raw)
            except Exception as e:
              rec={"condition":cond,"instance_id":iid,"model":MODEL,"provider":"Groq","stage":"FINAL_DECISION","status":"MALFORMED","raw_response":raw,"error":f"JSON parse failure: {e}","attempt":attempt,"latency_seconds":round(lat,4),"timestamp_utc":datetime.now(timezone.utc).isoformat()};break
            ok,err=validate_payload(payload)
            if not ok:
              rec={"condition":cond,"instance_id":iid,"model":MODEL,"provider":"Groq","stage":"FINAL_DECISION","status":"MALFORMED","raw_response":raw,"error":err,"attempt":attempt,"latency_seconds":round(lat,4),"timestamp_utc":datetime.now(timezone.utc).isoformat()};break
            u=getattr(resp,"usage",None)
            rec={"condition":cond,"instance_id":iid,"model":MODEL,"provider":"Groq","stage":"FINAL_DECISION","status":"SUCCESS",**payload,"raw_response":raw,"attempt":attempt,"latency_seconds":round(lat,4),"prompt_tokens":getattr(u,"prompt_tokens",None) if u else None,"completion_tokens":getattr(u,"completion_tokens",None) if u else None,"total_tokens":getattr(u,"total_tokens",None) if u else None,"timestamp_utc":datetime.now(timezone.utc).isoformat()};break
          except Exception as e:
            lat=time.perf_counter()-st; msg=str(e); rate="429" in msg or "rate_limit" in msg.lower() or "rate limit" in msg.lower()
            if attempt<MAX_RETRIES:
              wait=(retry_after(msg) if rate else TECHNICAL_RETRY_SECONDS)
              if wait is None: wait=60
              if rate: wait+=2
              print(f"[{idx:03d}/480] {'RATE_LIMIT' if rate else 'TECH_RETRY'} {cond} {iid} attempt {attempt}; sleeping {wait:.1f}s")
              time.sleep(wait);continue
            rec={"condition":cond,"instance_id":iid,"model":MODEL,"provider":"Groq","stage":"FINAL_DECISION","status":"API_ERROR","error":msg,"attempt":attempt,"latency_seconds":round(lat,4),"timestamp_utc":datetime.now(timezone.utc).isoformat()};break
        out.write(json.dumps(rec,ensure_ascii=False)+"\n");out.flush()
        suffix=f" -> {rec['decision']} ({rec['confidence']})" if rec.get("status")=="SUCCESS" else ""
        print(f"[{idx:03d}/480] {rec['status']:<10} {cond} {iid}{suffix}")
    print(f"\nFinished. Raw audit log: {outfile}")
if __name__=="__main__": main()
